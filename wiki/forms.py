"""
Forms for the wiki application.
"""

from pathlib import Path
import logging

from django import forms
from django.conf import settings
from django.contrib.auth.forms import AuthenticationForm, UserCreationForm
from django.contrib.auth import get_user_model
from django.forms import inlineformset_factory
from martor.fields import MartorFormField
from martor.widgets import MartorWidget

from .models import (
    Article,
    Category,
    Choice,
    CodingExercise,
    CodingTestCase,
    Comment,
    Profile,
    Question,
    UploadedFile,
)

User = get_user_model()

logger = logging.getLogger(__name__)


def extract_quiz_text(upload):
    """Extract text from a PDF, DOCX, or TXT upload."""
    file_ext = Path(upload.name).suffix.lower()
    upload.seek(0)

    if file_ext == ".txt":
        return upload.read().decode("utf-8", errors="ignore")

    if file_ext == ".docx":
        from docx import Document

        document = Document(upload)
        return "\n".join(paragraph.text for paragraph in document.paragraphs)

    if file_ext == ".pdf":
        from PyPDF2 import PdfReader

        reader = PdfReader(upload)
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    return ""


class SignUpForm(UserCreationForm):
    """Form for user registration."""

    email = forms.EmailField(required=True)

    class Meta(UserCreationForm.Meta):
        """Metadata for the SignUpForm."""

        model = User
        fields = ("username", "first_name", "last_name", "email")

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        placeholders = {
            "username": "Tên đăng nhập",
            "first_name": "Tên",
            "last_name": "Họ",
            "email": "Email",
            "password1": "Mật khẩu",
            "password2": "Nhập lại mật khẩu",
        }
        for name, field in self.fields.items():
            field.widget.attrs.update(
                {
                    "class": "form-control form-control-lg",
                    "placeholder": placeholders.get(name, ""),
                }
            )


class LoginForm(AuthenticationForm):
    """Form for user login."""

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        for field in self.fields.values():
            field.widget.attrs.update({"class": "form-control"})


class UserUpdateForm(forms.ModelForm):
    """Form for updating basic user information."""

    class Meta:
        """Metadata for the UserUpdateForm."""

        model = User
        fields = ("first_name", "last_name", "email")

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        for field in self.fields.values():
            field.widget.attrs.update({"class": "form-control"})


class ProfileForm(forms.ModelForm):
    """Form for updating user profile."""

    class Meta:
        """Metadata for the ProfileForm."""

        model = Profile
        fields = (
            "display_name",
            "bio",
            "avatar",
            "is_profile_private",
            "show_email_publicly",
        )
        widgets = {
            "bio": forms.Textarea(
                attrs={
                    "rows": 4,
                    "class": "form-control",
                    "placeholder": "Chia sẻ một chút về lĩnh vực bạn đang học...",
                }
            ),
            "avatar": forms.FileInput(attrs={"class": "form-control"}),
            "is_profile_private": forms.CheckboxInput(
                attrs={"class": "form-check-input"}
            ),
            "show_email_publicly": forms.CheckboxInput(
                attrs={"class": "form-check-input"}
            ),
        }

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.fields["display_name"].widget.attrs.update(
            {
                "class": "form-control",
                "placeholder": "Tên hiển thị",
            }
        )
        self.fields["is_profile_private"].help_text = (
            "Khi bật, người khác sẽ chỉ thấy trạng thái hồ sơ riêng tư."
        )
        self.fields["show_email_publicly"].help_text = (
            "Email chỉ hiển thị cho người khác khi hồ sơ đang ở chế độ công khai."
        )


class CategoryForm(forms.ModelForm):
    """Form for creating/updating categories."""

    class Meta:
        """Metadata for the CategoryForm."""

        model = Category
        fields = ("name", "slug", "description")
        widgets = {
            "description": forms.Textarea(
                attrs={
                    "rows": 4,
                    "class": "form-control",
                    "placeholder": "Mô tả ngắn cho danh mục",
                }
            ),
        }

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.fields["name"].widget.attrs.update(
            {
                "class": "form-control",
                "placeholder": "Ví dụ: Thuật toán đồ thị",
            }
        )
        self.fields["slug"].widget.attrs.update(
            {
                "class": "form-control",
                "placeholder": "thuat-toan-do-thi",
            }
        )


class ArticleForm(forms.ModelForm):
    """Form for creating/updating articles."""

    change_summary = forms.CharField(
        max_length=255,
        required=False,
        widget=forms.TextInput(
            attrs={
                "class": "form-control",
                "placeholder": "Tóm tắt các thay đổi của bạn (không bắt buộc)...",
            }
        ),
    )
    content = MartorFormField(
        widget=MartorWidget(
            attrs={
                "class": "form-control",
                "placeholder": (
                    "# Tiêu đề ghi chú\n\n" "Mô tả ý tưởng, độ phức tạp, ví dụ code..."
                ),
            }
        )
    )

    class Meta:
        """Metadata for the ArticleForm."""

        model = Article
        fields = (
            "title",
            "slug",
            "category",
            "tags",
            "allow_comments",
            "content",
            "change_summary",
        )

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.fields["tags"].widget.attrs.update(
            {
                "class": "form-control",
                "placeholder": "Gắn thẻ (ví dụ: segment-tree, dp, math)",
            }
        )
        self.fields["tags"].help_text = "Các thẻ cách nhau bằng dấu phẩy."
        self.fields["title"].widget.attrs.update(
            {
                "class": "form-control form-control-lg",
                "placeholder": "Ví dụ: Segment Tree cơ bản",
            }
        )
        self.fields["slug"].widget.attrs.update(
            {
                "class": "form-control",
                "placeholder": "segment-tree-co-ban",
            }
        )
        self.fields["slug"].help_text = (
            "Có thể để trống, hệ thống sẽ tự tạo từ tiêu đề."
        )
        self.fields["category"].widget.attrs.update({"class": "form-select"})
        self.fields["category"].required = False
        self.fields["category"].empty_label = "Chưa phân loại"
        self.fields["category"].help_text = (
            "Có thể để trống, bài viết sẽ tự chuyển vào danh mục Chưa phân loại."
        )
        self.fields["allow_comments"].widget.attrs.update({"class": "form-check-input"})
        self.fields["allow_comments"].help_text = (
            "Bật để người dùng có quyền được phép bình luận dưới bài viết này."
        )


class CommentForm(forms.ModelForm):
    """Form for adding comments."""

    content = MartorFormField(
        widget=MartorWidget(
            attrs={
                "class": "form-control",
                "placeholder": (
                    "Đặt câu hỏi, góp ý lời giải, hoặc bổ sung mẹo tối ưu..."
                ),
            }
        )
    )

    class Meta:
        """Metadata for the CommentForm."""

        model = Comment
        fields = ("content",)


class UploadFileForm(forms.ModelForm):
    """Form for uploading files."""

    class Meta:
        """Metadata for the UploadFileForm."""

        model = UploadedFile
        fields = ("file", "description")
        widgets = {
            "description": forms.TextInput(
                attrs={
                    "class": "form-control",
                    "placeholder": "Mô tả ngắn file (tùy chọn)",
                }
            ),
        }

    def clean_file(self):
        """Validate uploaded file size."""
        file = self.cleaned_data.get("file")
        if file and file.size > 15 * 1024 * 1024:
            raise forms.ValidationError("File quá lớn. Vui lòng upload tối đa 15MB.")
        return file


class QuestionForm(forms.ModelForm):
    """Form for creating quiz questions."""

    question_file = forms.FileField(
        label="Upload nội dung từ file",
        required=False,
        help_text="Có thể upload PDF/DOCX/TXT. Nếu để trống nội dung, hệ thống sẽ lấy text từ file.",
        widget=forms.FileInput(
            attrs={
                "class": "form-control",
                "accept": ".pdf,.docx,.txt",
            }
        ),
    )
    content = MartorFormField(
        label="Nội dung câu hỏi",
        required=False,
        widget=MartorWidget(
            attrs={
                "class": "form-control",
                "placeholder": "Nội dung câu hỏi (hỗ trợ Markdown)...",
            }
        ),
    )
    explanation = MartorFormField(
        label="Giải thích đáp án",
        required=False,
        widget=MartorWidget(
            attrs={
                "class": "form-control",
                "placeholder": (
                    "Giải thích đáp án (sẽ hiển thị sau khi người dùng trả lời)..."
                ),
            }
        ),
    )

    class Meta:
        """Metadata for the QuestionForm."""

        model = Question
        fields = ("content", "question_file", "explanation", "order")
        widgets = {
            "order": forms.NumberInput(
                attrs={
                    "class": "form-control",
                    "placeholder": "Thứ tự ưu tiên hiển thị",
                }
            ),
        }

    def clean_question_file(self):
        """Validate uploaded source file for a quiz question."""
        upload = self.cleaned_data.get("question_file")
        if not upload:
            return upload

        allowed_extensions = {".pdf", ".docx", ".txt"}
        file_ext = Path(upload.name).suffix.lower()
        if file_ext not in allowed_extensions:
            raise forms.ValidationError("Chỉ hỗ trợ file PDF, DOCX hoặc TXT.")
        if upload.size > 15 * 1024 * 1024:
            raise forms.ValidationError("File quá lớn. Vui lòng upload tối đa 15MB.")
        return upload

    def clean(self):
        """Allow typing content directly or extracting it from an uploaded file."""
        cleaned_data = super().clean()
        content = (cleaned_data.get("content") or "").strip()
        upload = cleaned_data.get("question_file")

        if not content and upload:
            try:
                extracted_text = extract_quiz_text(upload).strip()
            except Exception as error:  # pylint: disable=broad-exception-caught
                logger.exception("Failed to extract quiz text from uploaded file")
                self.add_error(
                    "question_file",
                    f"Không thể đọc file này: {error}",
                )
                return cleaned_data
            if extracted_text:
                cleaned_data["content"] = extracted_text
            else:
                self.add_error(
                    "question_file",
                    "Không trích xuất được nội dung từ file này.",
                )
        elif not content:
            self.add_error("content", "Nhập nội dung câu hỏi hoặc upload file.")

        return cleaned_data


class ChoiceForm(forms.ModelForm):
    """Form for quiz choices."""

    class Meta:
        """Metadata for the ChoiceForm."""

        model = Choice
        fields = ("content", "is_correct")
        widgets = {
            "content": forms.TextInput(
                attrs={
                    "class": "form-control",
                    "placeholder": "Nhập nội dung lựa chọn...",
                }
            ),
            "is_correct": forms.CheckboxInput(attrs={"class": "form-check-input"}),
        }


def get_code_language_choices():
    """Expose enabled judge languages as form choices."""
    from .services.code_runner import get_enabled_language_choices
    return [(c["key"], c["label"]) for c in get_enabled_language_choices()]


class CodingExerciseForm(forms.ModelForm):
    """Form for configuring a coding exercise."""

    allowed_languages = forms.MultipleChoiceField(
        choices=(),
        required=False,
        widget=forms.CheckboxSelectMultiple,
    )

    checker_language = forms.ChoiceField(
        choices=[
            ("python", "Python 3"),
            ("cpp", "C++17"),
        ],
        required=False,
        widget=forms.Select(attrs={"class": "form-select"}),
        label="Ngôn ngữ checker",
    )

    class Meta:
        """Metadata for CodingExerciseForm."""

        model = CodingExercise
        fields = (
            "title",
            "description",
            "is_enabled",
            "allowed_languages",
            "default_language",
            "time_limit_ms",
            "memory_limit_mb",
            "compare_mode",
            "checker_code",
            "checker_language",
        )
        labels = {
            "compare_mode": "Checker",
            "time_limit_ms": "Time limit (ms)",
            "memory_limit_mb": "Memory limit (MB)",
            "checker_code": "Mã nguồn checker",
            "checker_language": "Ngôn ngữ checker",
        }
        widgets = {
            "description": forms.Textarea(
                attrs={
                    "rows": 5,
                    "class": "form-control",
                    "placeholder": "Mô tả bài tập, yêu cầu input/output...",
                }
            ),
            "is_enabled": forms.CheckboxInput(attrs={"class": "form-check-input"}),
            "default_language": forms.Select(attrs={"class": "form-select"}),
            "time_limit_ms": forms.NumberInput(attrs={"class": "form-control"}),
            "memory_limit_mb": forms.NumberInput(attrs={"class": "form-control"}),
            "compare_mode": forms.Select(attrs={"class": "form-select"}),
            "checker_code": forms.Textarea(
                attrs={
                    "rows": 10,
                    "class": "form-control font-monospace",
                    "placeholder": "Nhập mã nguồn checker tại đây...\n\nĐối với Python/C++, checker sẽ nhận 3 đối số dòng lệnh:\n1. Đường dẫn file input (stdin)\n2. Đường dẫn file output của user (actual output)\n3. Đường dẫn file output mẫu (expected output)\n\nTrả về exit code 0 nếu chấp nhận (AC), trả về exit code 1 hoặc 2 nếu sai (WA).",
                }
            ),
        }

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        language_choices = get_code_language_choices()
        self.fields["allowed_languages"].choices = language_choices
        self.fields["allowed_languages"].help_text = (
            "Chỉ các ngôn ngữ được chọn mới hiện ra ở Monaco editor."
        )
        self.fields["default_language"].choices = [
            ("", "-- Chọn ngôn ngữ mặc định --")
        ] + language_choices
        self.fields["title"].widget.attrs.update(
            {"class": "form-control", "placeholder": "VD: In ra tổng 2 số"}
        )

        starter_map = self.instance.starter_code_map if self.instance.pk else {}
        from .services.code_runner import _merged_configs
        language_configs = _merged_configs()
        for language_key, config in language_configs.items():
            field_name = f"starter_code_{language_key}"
            self.fields[field_name] = forms.CharField(
                required=False,
                label=f"Starter code - {config.get('label', language_key)}",
                widget=forms.Textarea(
                    attrs={
                        "rows": 8,
                        "class": "form-control font-monospace",
                        "placeholder": "Code mẫu cho ngôn ngữ này...",
                    }
                ),
                initial=starter_map.get(
                    language_key, config.get("starter_code", "")
                ),
            )

    def clean(self):
        """Ensure default language is part of the allowed set."""
        cleaned_data = super().clean()
        allowed_languages = cleaned_data.get("allowed_languages") or []
        default_language = cleaned_data.get("default_language")
        if default_language and default_language not in allowed_languages:
            self.add_error(
                "default_language", "Ngôn ngữ mặc định phải nằm trong danh sách cho phép."
            )
        if cleaned_data.get("is_enabled") and not allowed_languages:
            self.add_error(
                "allowed_languages", "Cần chọn ít nhất một ngôn ngữ khi bật bài tập code."
            )
        if cleaned_data.get("compare_mode") == "custom_checker":
            checker_code = cleaned_data.get("checker_code")
            checker_language = cleaned_data.get("checker_language")
            if not checker_code:
                self.add_error("checker_code", "Vui lòng nhập mã nguồn checker.")
            if not checker_language:
                self.add_error("checker_language", "Vui lòng chọn ngôn ngữ checker.")
        return cleaned_data

    def save(self, commit=True):
        """Persist starter codes as a JSON map."""
        instance = super().save(commit=False)
        starter_map = {}
        from .services.code_runner import _merged_configs
        for language_key in _merged_configs():
            starter_value = self.cleaned_data.get(f"starter_code_{language_key}", "")
            if starter_value:
                starter_map[language_key] = starter_value
        instance.starter_code_map = starter_map
        if commit:
            instance.save()
            self.save_m2m()
        return instance


class CodingTestCaseForm(forms.ModelForm):
    """Form for creating and editing coding exercise testcases."""

    class Meta:
        """Metadata for CodingTestCaseForm."""

        model = CodingTestCase
        fields = (
            "name",
            "input_text",
            "expected_output_text",
            "input_file",
            "expected_output_file",
            "is_sample",
            "order",
            "score",
            "subtask_id",
        )
        labels = {
            "name": "Tên testcase",
            "input_text": "Dữ liệu đầu vào (Input text)",
            "expected_output_text": "Dữ liệu đầu ra mong đợi (Expected Output text)",
            "input_file": "Tệp dữ liệu đầu vào (Input file)",
            "expected_output_file": "Tệp dữ liệu đầu ra mong đợi (Output file)",
            "is_sample": "Là testcase mẫu (Sample testcase)",
            "order": "Thứ tự thực hiện (Order)",
            "score": "Điểm số (Score)",
            "subtask_id": "ID nhóm bài phụ (Subtask ID)",
        }
        help_texts = {
            "name": "Tên định danh cho testcase (Ví dụ: sample-1, tc-01).",
            "input_text": "Nội dung đầu vào dạng văn bản. Để trống nếu bạn tải tệp lên.",
            "expected_output_text": "Nội dung đầu ra kỳ vọng dạng văn bản. Để trống nếu bạn tải tệp lên.",
            "input_file": "Tải lên tệp input (.inp, .in, .txt). Sẽ được ưu tiên hơn phần văn bản nếu nhập cả hai.",
            "expected_output_file": "Tải lên tệp output/expected (.out, .ans, .txt). Sẽ được ưu tiên hơn phần văn bản nếu nhập cả hai.",
            "is_sample": "Tích chọn nếu muốn hiển thị công khai testcase này làm mẫu ví dụ trên bài viết.",
            "order": "Thứ tự sắp xếp và đánh giá của testcase (số nhỏ chạy trước).",
            "score": "Số điểm đạt được khi vượt qua testcase này.",
            "subtask_id": "ID nhóm chấm bài phụ. Các testcase có cùng ID sẽ gom lại thành 1 subtask chấm điểm theo cụm.",
        }
        widgets = {
            "name": forms.TextInput(
                attrs={"class": "form-control", "placeholder": "Ví dụ: sample-1, test-01"}
            ),
            "input_text": forms.Textarea(
                attrs={
                    "rows": 6,
                    "class": "form-control font-monospace",
                    "placeholder": "Nhập dữ liệu đầu vào trực tiếp tại đây nếu không upload file...",
                }
            ),
            "expected_output_text": forms.Textarea(
                attrs={
                    "rows": 6,
                    "class": "form-control font-monospace",
                    "placeholder": "Nhập dữ liệu mong đợi đầu ra trực tiếp tại đây nếu không upload file...",
                }
            ),
            "input_file": forms.FileInput(attrs={"class": "form-control"}),
            "expected_output_file": forms.FileInput(attrs={"class": "form-control"}),
            "is_sample": forms.CheckboxInput(attrs={"class": "form-check-input"}),
            "order": forms.NumberInput(attrs={"class": "form-control", "placeholder": "Ví dụ: 1"}),
            "score": forms.NumberInput(attrs={"class": "form-control", "placeholder": "Ví dụ: 10"}),
            "subtask_id": forms.NumberInput(attrs={"class": "form-control", "placeholder": "Ví dụ: 1"}),
        }

    def clean(self):
        """Require either raw text or uploaded files for both sides."""
        cleaned_data = super().clean()
        if not (cleaned_data.get("input_text") or cleaned_data.get("input_file")):
            self.add_error("input_text", "Can input raw hoac file input.")
        if not (
            cleaned_data.get("expected_output_text")
            or cleaned_data.get("expected_output_file")
        ):
            self.add_error(
                "expected_output_text", "Can output raw hoac file output."
            )
        return cleaned_data


class BaseChoiceFormSet(forms.models.BaseInlineFormSet):
    def clean(self):
        super().clean()
        if any(self.errors):
            return

        completed_forms = 0
        correct_count = 0
        for form in self.forms:
            if not form.cleaned_data or form.cleaned_data.get('DELETE'):
                continue
            completed_forms += 1
            if form.cleaned_data.get('is_correct'):
                correct_count += 1

        has_file = bool(self.files and self.files.get("question_file"))
        if completed_forms > 0:
            if completed_forms < 2:
                raise forms.ValidationError("Một câu hỏi cần có ít nhất 2 lựa chọn.")
            if correct_count == 0:
                raise forms.ValidationError("Cần chọn ít nhất một lựa chọn làm đáp án đúng.")
            if correct_count > 1:
                raise forms.ValidationError("Chỉ được chọn duy nhất một đáp án đúng.")
        elif not has_file:
            raise forms.ValidationError("Một câu hỏi cần có ít nhất 2 lựa chọn.")


ChoiceFormSet = inlineformset_factory(
    Question, Choice, form=ChoiceForm, formset=BaseChoiceFormSet, extra=4, can_delete=True
)
